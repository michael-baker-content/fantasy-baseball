import csv
import json
from datetime import date

import pytest
import export_mlb_ytd as exporter

from export_mlb_ytd import (
    BIO_FIELDS, HITTING_FIELDS, PITCHING_FIELDS, PLAYOFF_TEAMS, build_rows,
    is_injured, is_optioned, write_excel,
)


@pytest.fixture(autouse=True)
def block_live_api(monkeypatch):
    def fail(*args, **kwargs):
        raise AssertionError("Exporter tests must not access the live MLB API")
    monkeypatch.setattr(exporter, "api_get", fail)


def test_roster_status_flags():
    assert is_optioned("MIN", "Minors")
    assert not is_optioned("A", "Active")
    assert is_injured("IL10", "10-Day Injured List")
    assert not is_injured("A", "Active")


def test_build_rows_excludes_unaffiliated_players():
    splits = [
        {"player": {"id": 1}, "stat": {"gamesPlayed": 2, "hits": 3}},
        {"player": {"id": 2}, "stat": {"gamesPlayed": 4, "hits": 5}},
    ]
    people = {
        1: {"id": 1, "fullName": "Affiliated Player"},
        2: {"id": 2, "fullName": "Free Agent"},
    }
    affiliations = {
        1: {"organization_name": "Test Club", "status_code": "A", "status": "Active"},
    }
    rows = build_rows(splits, {"hits": "hits"}, affiliations, people)
    assert [row["player_name"] for row in rows] == ["Affiliated Player"]
    assert rows[0]["hits"] == 3


def test_hitting_rows_require_an_offensive_appearance():
    splits = [
        {"player": {"id": 1}, "stat": {"gamesPlayed": 3, "plateAppearances": 0}},
        {"player": {"id": 2}, "stat": {"gamesPlayed": 1, "plateAppearances": 0, "runs": 1}},
    ]
    people = {player_id: {"id": player_id, "fullName": f"Player {player_id}"} for player_id in (1, 2)}
    affiliations = {
        player_id: {"organization_name": "Test Club", "status_code": "A", "status": "Active"}
        for player_id in (1, 2)
    }
    rows = build_rows(splits, {"runs": "runs"}, affiliations, people, "hitting")
    assert [row["player_id"] for row in rows] == [2]


def test_playoff_pool_has_requested_teams():
    assert len(PLAYOFF_TEAMS) == 14
    assert "Arizona Diamondbacks" in PLAYOFF_TEAMS
    assert "Texas Rangers" in PLAYOFF_TEAMS


def test_excel_output_has_batters_and_pitchers_sheets(tmp_path):
    path = tmp_path / "mlb_players_ytd_2026-09-29.xlsx"
    write_excel(path, [], [])
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook.sheetnames == ["Batters", "Pitchers"]
        assert [cell.value for cell in workbook["Batters"][1]] == BIO_FIELDS + list(HITTING_FIELDS)
        assert [cell.value for cell in workbook["Pitchers"][1]] == BIO_FIELDS + list(PITCHING_FIELDS)
    finally:
        workbook.close()


def test_current_affiliations_retains_minors_and_il_excludes_released_and_fa(monkeypatch):
    calls = []

    def fake_api(path, params):
        calls.append((path, params))
        if path == "/teams":
            return {"teams": [{"id": 10, "name": "Test Club"}]}
        statuses = [("A", "Active"), ("MIN", "Minors"), ("IL60", "60-Day Injured List"),
                    ("RL", "Released"), ("FA", "Free Agent"), ("X", "Released"),
                    ("X", "Free Agency")]
        return {"roster": [
            {"person": {"id": index}, "status": {"code": code, "description": description}}
            for index, (code, description) in enumerate(statuses, 1)
        ]}

    monkeypatch.setattr(exporter, "api_get", fake_api)
    affiliations, _ = exporter.current_affiliations("2026-09-21", 2026, {"Test Club"})
    assert set(affiliations) == {1, 2, 3}
    assert calls[1][1] == {"rosterType": "fullRoster", "date": "2026-09-21"}


@pytest.mark.parametrize("output_format", ["csv", "xlsx", "both"])
def test_range_export_uses_current_affiliation_and_all_team_totals(monkeypatch, tmp_path, output_format):
    class FixedDate(date):
        @classmethod
        def today(cls):
            return cls(2026, 9, 21)

    calls = []

    def fake_api(path, params):
        calls.append((path, params))
        if path == "/teams":
            assert params["season"] == 2026
            return {"teams": [{"id": 10, "name": "Test Club"}]}
        if path == "/teams/10/roster":
            assert params["date"] == "2026-09-21"
            return {"roster": [{"person": {"id": 1},
                                "status": {"code": "MIN", "description": "Minors"}}]}
        if path in {"/stats", "/people/1/stats"}:
            assert params["startDate"] == "05/01/2026"
            assert params["endDate"] == "05/30/2026"
            assert params["gameType"] == "R"
            assert "teamId" not in params
            if params["group"] == "pitching":
                return {"stats": [{"splits": []}]}
            if path == "/people/1/stats":
                # Official combined totals: preserve rates rather than adding them.
                splits = [{"stat": {"gamesPlayed": 20, "plateAppearances": 80,
                                     "hits": 21, "atBats": 70, "avg": ".300"}}]
            else:
                # Both stints predate the player's move to the selected club.
                splits = [
                    {"player": {"id": 1}, "team": {"id": 20},
                     "stat": {"gamesPlayed": 10, "hits": 9, "atBats": 30}},
                    {"player": {"id": 1}, "team": {"id": 30},
                     "stat": {"gamesPlayed": 10, "hits": 12, "atBats": 40}},
                    {"player": {"id": 2}, "team": {"id": 10},
                     "stat": {"gamesPlayed": 10, "hits": 15}},
                ]
            return {"stats": [{"splits": splits}]}
        if path == "/people":
            assert params["personIds"] == "1"
            return {"people": [{"id": 1, "fullName": "Traded Player", "primaryPosition": {"abbreviation": "SS"}}]}
        raise AssertionError(f"Unexpected API request: {path}")

    monkeypatch.setattr(exporter, "date", FixedDate)
    monkeypatch.setattr(exporter, "ROOT", tmp_path)
    monkeypatch.setattr(exporter, "PLAYOFF_TEAMS", {"Test Club"})
    monkeypatch.setattr(exporter, "api_get", fake_api)
    monkeypatch.setattr("sys.argv", ["export_mlb_ytd.py", "--season", "2026",
                                    "--start", "2026-05-01", "--through", "2026-05-30",
                                    "--format", output_format, "--site"])
    assert exporter.main() == 0
    site = json.loads((tmp_path / "docs/data/player-stats.json").read_text(encoding="utf-8"))
    published = site["ranges"][0]
    assert published["start"] == "2026-05-01"
    assert published["through"] == "2026-05-30"
    assert published["roster_date"] == "2026-09-21"
    assert len(published["batters"]) == 1
    assert published["batters"][0]["team"] == "Test Club"
    assert published["batters"][0]["position"] == "SS"
    assert published["batters"][0]["stats"]["H"] == 21
    assert published["pitchers"] == []
    output = tmp_path / "data" / "2026"
    assert len(list(output.glob("*.csv"))) == (2 if output_format in {"csv", "both"} else 0)
    assert len(list(output.glob("*.xlsx"))) == (1 if output_format in {"xlsx", "both"} else 0)
    if output_format in {"csv", "both"}:
        with (output / "mlb_hitters_2026-05-01_through_2026-05-30.csv").open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        assert len(rows) == 1
        assert rows[0]["current_mlb_organization"] == "Test Club"
        assert rows[0]["optioned_to_minors"] == "True"
        assert rows[0]["hits"] == "21"
        assert rows[0]["batting_average"] == ".300"
    if output_format in {"xlsx", "both"}:
        from openpyxl import load_workbook
        workbook = load_workbook(output / "mlb_players_2026-05-01_through_2026-05-30.xlsx")
        try:
            assert workbook.sheetnames == ["Batters", "Pitchers"]
            fields = BIO_FIELDS + list(HITTING_FIELDS)
            values = [cell.value for cell in workbook["Batters"][2]]
            row = dict(zip(fields, values))
            assert row["current_mlb_organization"] == "Test Club"
            assert row["hits"] == 21
            assert row["batting_average"] == .3
            assert "2026-05-01 through 2026-05-30" in workbook.properties.title
            assert "2026-09-21" in workbook.properties.description
        finally:
            workbook.close()


def test_combined_splits_prefers_total_without_double_counting():
    splits = [{"player": {"id": 1}, "team": {"id": 20}, "stat": {"hits": 4}},
              {"player": {"id": 1}, "team": {"id": 30}, "stat": {"hits": 5}},
              {"player": {"id": 1}, "isTotal": True, "stat": {"hits": 9}}]
    rows = exporter.combined_splits(splits, "hitting", 2026, "2026-05-01", "2026-05-30", {1})
    assert len(rows) == 1
    assert rows[0]["stat"]["hits"] == 9


def test_combined_splits_rejects_ambiguous_totals(monkeypatch):
    splits = [{"player": {"id": 1}, "team": {"id": team}, "stat": {"hits": 4}}
              for team in (20, 30)]
    monkeypatch.setattr(exporter, "api_get", lambda *args: {"stats": [{"splits": splits}]})
    with pytest.raises(ValueError, match="Cannot determine combined hitting totals"):
        exporter.combined_splits(splits, "hitting", 2026, "2026-05-01", "2026-05-30", {1})


@pytest.mark.parametrize("start,through", [
    ("2026-05-31", "2026-05-01"), ("2025-05-01", "2026-05-30"),
    ("2026-05-01", "2027-05-30"), ("2026-02-30", "2026-05-30"),
])
def test_invalid_dates_fail_before_network(monkeypatch, start, through):
    monkeypatch.setattr("sys.argv", ["export_mlb_ytd.py", "--season", "2026",
                                    "--start", start, "--through", through])
    with pytest.raises(SystemExit) as error:
        exporter.main()
    assert error.value.code == 2


def test_default_start_is_january_first(monkeypatch):
    def fake_api(path, params):
        assert params["startDate"] == "01/01/2026"
        assert params["endDate"] == "05/30/2026"
        return {"stats": [{"splits": []}]}

    monkeypatch.setattr(exporter, "api_get", fake_api)
    assert exporter.season_splits("hitting", 2026, "2026-05-30") == []


def test_site_ranges_replace_ytd_preserve_custom_and_use_selected_columns(tmp_path):
    path = tmp_path / "player-stats.json"
    batter = {"player_id": 1, "player_name": "Test Batter", "current_mlb_organization": "Test Club",
              "games_played": 30, "at_bats": 100, "hits": 30, "runs": 20, "home_runs": 5,
              "rbi": 15, "stolen_bases": 3, "walks": 10, "batting_average": ".300",
              "birth_date": "2000-01-01", "primary_position": "C"}
    pitcher = {"player_id": 2, "player_name": "Test Pitcher", "current_mlb_organization": "Test Club",
               "games": 6, "games_started": 4, "innings_pitched": "30.2", "wins": 3, "losses": 1, "saves": 0,
               "strikeouts": 40, "era": "2.35", "whip": "1.01", "primary_position": "P"}
    for start, through in [("2026-01-01", "2026-09-20"), ("2026-05-01", "2026-05-30"),
                           ("2026-01-01", "2026-09-21"), ("2026-05-01", "2026-05-30")]:
        exporter.write_site_data(path, [batter], [pitcher], 2026, start, through, "2026-09-21")
    ranges = json.loads(path.read_text(encoding="utf-8"))["ranges"]
    assert len(ranges) == 2
    assert ranges[0]["id"] == "ytd"
    assert ranges[0]["through"] == "2026-09-21"
    assert ranges[1]["start"] == "2026-05-01"
    assert set(ranges[0]["batters"][0]["stats"]) == {"G", "AB", "H", "R", "HR", "RBI", "SB", "BB", "AVG"}
    assert set(ranges[0]["pitchers"][0]["stats"]) == {"G", "GS", "IP", "W", "L", "SV", "K", "ERA", "WHIP"}
    assert ranges[0]["pitchers"][0]["stats"]["GS"] == 4
    assert ranges[0]["pitchers"][0]["stats"]["IP"] == "30.2"
    assert ranges[0]["batters"][0]["position"] == "C"
    assert ranges[0]["pitchers"][0]["position"] == "P"
    assert "birth_date" not in ranges[0]["batters"][0]
    exporter.write_site_data(path, [], [], 2027, "2027-01-01", "2027-04-01", "2027-04-01")
    ranges = json.loads(path.read_text(encoding="utf-8"))["ranges"]
    assert len(ranges) == 1
    assert ranges[0]["season"] == 2027


@pytest.mark.parametrize("through,expected", [
    ("2026-09-21", "2026-08-23"),
    ("2026-03-01", "2026-01-31"),
    ("2024-03-01", "2024-02-01"),
    ("2026-01-15", "2026-01-01"),
])
def test_last_30_days_are_inclusive_and_bounded_to_season(through, expected):
    end = date.fromisoformat(through)
    assert exporter.last_30_start(end.year, end).isoformat() == expected


def test_last_30_cli_publishes_range_and_preserves_ytd(monkeypatch, tmp_path):
    path = tmp_path / "docs/data/player-stats.json"
    exporter.write_site_data(path, [], [], 2026, "2026-01-01", "2026-09-21", "2026-09-21")
    exporter.write_site_data(path, [], [], 2026, "2026-08-22", "2026-09-20", "2026-09-20", last_30_days=True)
    calls = []

    def fake_splits(group, season, through, start):
        calls.append((group, season, through, start))
        return []

    monkeypatch.setattr(exporter, "ROOT", tmp_path)
    monkeypatch.setattr(exporter, "current_affiliations", lambda *args: ({}, {}))
    monkeypatch.setattr(exporter, "season_splits", fake_splits)
    monkeypatch.setattr("sys.argv", ["export_mlb_ytd.py", "--season", "2026", "--through", "2026-09-21",
                                    "--last-30-days", "--format", "csv", "--site"])
    assert exporter.main() == 0
    assert calls == [(group, 2026, "2026-09-21", "2026-08-23") for group in ("hitting", "pitching")]
    ranges = json.loads(path.read_text(encoding="utf-8"))["ranges"]
    assert [row["id"] for row in ranges] == ["ytd", "last30"]
    assert ranges[1]["start"] == "2026-08-23"
    assert ranges[1]["through"] == "2026-09-21"
    assert ranges[1]["label"] == "Last 30 days"
    assert (tmp_path / "data/2026/mlb_pitchers_last30_2026-08-23_through_2026-09-21.csv").exists()


def test_last_30_days_cannot_be_combined_with_explicit_start(monkeypatch):
    monkeypatch.setattr("sys.argv", ["export_mlb_ytd.py", "--start", "2026-05-01", "--last-30-days"])
    with pytest.raises(SystemExit) as error:
        exporter.main()
    assert error.value.code == 2
