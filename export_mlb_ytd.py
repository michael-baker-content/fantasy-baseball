"""Export regular-season date-range stats for the current MLB playoff pool."""

from __future__ import annotations

import argparse
import csv
import json
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen


BASE_URL = "https://statsapi.mlb.com/api/v1"
ROOT = Path(__file__).resolve().parent

PLAYOFF_TEAMS = {
    "Arizona Diamondbacks", "Atlanta Braves", "Boston Red Sox", "Chicago Cubs",
    "Chicago White Sox", "Cleveland Guardians", "Houston Astros",
    "Los Angeles Dodgers", "Milwaukee Brewers", "New York Yankees",
    "Philadelphia Phillies", "San Diego Padres", "Tampa Bay Rays", "Texas Rangers",
}

BIO_FIELDS = [
    "player_id", "player_name", "current_mlb_organization", "roster_status_code",
    "roster_status", "optioned_to_minors", "injured_list", "primary_position",
    "bats", "throws", "age", "birth_date", "birth_city", "birth_state_province",
    "birth_country", "nationality", "college", "draft_year", "mlb_debut_date",
]

HITTING_FIELDS = {
    "games_played": "gamesPlayed", "plate_appearances": "plateAppearances",
    "at_bats": "atBats", "runs": "runs", "hits": "hits", "doubles": "doubles",
    "triples": "triples", "home_runs": "homeRuns", "rbi": "rbi",
    "total_bases": "totalBases", "walks": "baseOnBalls",
    "intentional_walks": "intentionalWalks", "strikeouts": "strikeOuts",
    "hit_by_pitch": "hitByPitch", "stolen_bases": "stolenBases",
    "caught_stealing": "caughtStealing", "sacrifice_flies": "sacFlies",
    "sacrifice_bunts": "sacBunts", "grounded_into_double_plays": "groundIntoDoublePlay",
    "batting_average": "avg", "on_base_percentage": "obp",
    "slugging_percentage": "slg", "on_base_plus_slugging": "ops", "babip": "babip",
}

PITCHING_FIELDS = {
    "games": "gamesPlayed", "games_started": "gamesStarted",
    "games_finished": "gamesFinished", "innings_pitched": "inningsPitched",
    "batters_faced": "battersFaced", "wins": "wins", "losses": "losses",
    "win_percentage": "winPercentage", "saves": "saves",
    "save_opportunities": "saveOpportunities", "holds": "holds",
    "blown_saves": "blownSaves", "complete_games": "completeGames",
    "shutouts": "shutouts", "hits_allowed": "hits", "runs_allowed": "runs",
    "earned_runs": "earnedRuns", "home_runs_allowed": "homeRuns",
    "walks": "baseOnBalls", "intentional_walks": "intentionalWalks",
    "hit_batters": "hitBatsmen", "strikeouts": "strikeOuts",
    "wild_pitches": "wildPitches", "balks": "balks", "era": "era", "whip": "whip",
    "opponent_batting_average": "avg", "strikeouts_per_9": "strikeoutsPer9Inn",
    "walks_per_9": "walksPer9Inn", "hits_per_9": "hitsPer9Inn",
    "home_runs_per_9": "homeRunsPer9", "strikeout_walk_ratio": "strikeoutWalkRatio",
    "pitches": "numberOfPitches", "strikes": "strikes",
    "strike_percentage": "strikePercentage", "pitches_per_inning": "pitchesPerInning",
    "inherited_runners": "inheritedRunners",
    "inherited_runners_scored": "inheritedRunnersScored",
}


def api_get(path: str, params: dict | None = None) -> dict | list:
    url = f"{BASE_URL}{path}"
    if params:
        url += "?" + urlencode(params)
    request = Request(url, headers={"User-Agent": "BABBD-roto-export/1.0"})
    with urlopen(request, timeout=60) as response:
        return json.load(response)


def season_splits(group: str, season: int, through: str, start: str | None = None) -> list[dict]:
    payload = api_get("/stats", {
        "stats": "byDateRange", "group": group, "season": season, "gameType": "R",
        "startDate": datetime.strptime(start or f"{season}-01-01", "%Y-%m-%d").strftime("%m/%d/%Y"),
        "endDate": datetime.strptime(through, "%Y-%m-%d").strftime("%m/%d/%Y"),
        "sportIds": 1, "playerPool": "ALL", "limit": 5000,
    })
    return payload["stats"][0]["splits"]


def current_affiliations(as_of: str, season: int, selected_teams: set[str]) -> tuple[dict[int, dict], dict[int, dict]]:
    teams = api_get("/teams", {"sportId": 1, "season": season})["teams"]
    available = {team["name"] for team in teams}
    unknown = selected_teams - available
    if unknown:
        raise ValueError(f"Unknown MLB team name(s): {', '.join(sorted(unknown))}")
    teams = [team for team in teams if team["name"] in selected_teams]
    organizations = {team["id"]: team for team in teams}
    affiliations: dict[int, dict] = {}
    for team in teams:
        roster = api_get(f"/teams/{team['id']}/roster", {
            "rosterType": "fullRoster", "date": as_of,
        }).get("roster", [])
        for entry in roster:
            player_id = entry.get("person", {}).get("id")
            status_code = entry.get("status", {}).get("code", "")
            status = entry.get("status", {}).get("description", "")
            excluded = status_code.upper() in {"RL", "FA"} or any(
                word in status.lower() for word in ("released", "free agent", "free agency")
            )
            if player_id and not excluded:
                affiliations[player_id] = {
                    "organization_id": team["id"],
                    "organization_name": team["name"],
                    "status_code": status_code,
                    "status": status,
                }
    return affiliations, organizations


def people_details(player_ids: set[int]) -> dict[int, dict]:
    people: dict[int, dict] = {}
    ids = sorted(player_ids)
    for start in range(0, len(ids), 50):
        payload = api_get("/people", {
            "personIds": ",".join(map(str, ids[start:start + 50])),
            "hydrate": "education",
        })
        people.update({person["id"]: person for person in payload.get("people", [])})
    return people


def is_optioned(status_code: str, status: str) -> bool:
    value = f"{status_code} {status}".lower()
    return any(word in value for word in ("option", "minors", "minor league"))


def is_injured(status_code: str, status: str) -> bool:
    value = f"{status_code} {status}".lower()
    return status_code.upper().startswith("IL") or any(word in value for word in ("injur", "disabled"))


def bio_row(person: dict, affiliation: dict) -> dict:
    colleges = person.get("education", {}).get("colleges", [])
    return {
        "player_id": person.get("id", ""),
        "player_name": person.get("fullName", ""),
        "current_mlb_organization": affiliation["organization_name"],
        "roster_status_code": affiliation["status_code"],
        "roster_status": affiliation["status"],
        "optioned_to_minors": is_optioned(affiliation["status_code"], affiliation["status"]),
        "injured_list": is_injured(affiliation["status_code"], affiliation["status"]),
        "primary_position": person.get("primaryPosition", {}).get("abbreviation", ""),
        "bats": person.get("batSide", {}).get("code", ""),
        "throws": person.get("pitchHand", {}).get("code", ""),
        "age": person.get("currentAge", ""),
        "birth_date": person.get("birthDate", ""),
        "birth_city": person.get("birthCity", ""),
        "birth_state_province": person.get("birthStateProvince", ""),
        "birth_country": person.get("birthCountry", ""),
        "nationality": person.get("nationality", ""),
        "college": "; ".join(college.get("name", "") for college in colleges if college.get("name")),
        "draft_year": person.get("draftYear", ""),
        "mlb_debut_date": person.get("mlbDebutDate", ""),
    }


def build_rows(splits: list[dict], field_map: dict[str, str], affiliations: dict[int, dict],
               people: dict[int, dict], group: str = "") -> list[dict]:
    rows = []
    for split in splits:
        player_id = split.get("player", {}).get("id")
        if player_id not in affiliations or player_id not in people:
            continue
        stat = split.get("stat", {})
        if int(stat.get("gamesPlayed", 0) or 0) < 1:
            continue
        if group == "hitting" and not (
            int(stat.get("plateAppearances", 0) or 0) > 0
            or any(int(stat.get(field, 0) or 0) > 0 for field in ("runs", "stolenBases", "caughtStealing"))
        ):
            continue
        row = bio_row(people[player_id], affiliations[player_id])
        row.update({column: stat.get(api_name, "") for column, api_name in field_map.items()})
        rows.append(row)
    return sorted(rows, key=lambda row: (row["player_name"], row["player_id"]))


def combined_splits(splits: list[dict], group: str, season: int,
                    start: str, through: str, eligible_ids: set[int]) -> list[dict]:
    """Keep one all-team total per player, resolving team splits if necessary.

    Never add rate statistics or count an aggregate row alongside team rows.
    If the league endpoint supplies multiple rows, request that player's
    unfiltered date-range totals. Refuse ambiguous results instead of silently
    exporting duplicated or partial totals.
    """
    by_player: dict[int, list[dict]] = {}
    for split in splits:
        player_id = split.get("player", {}).get("id")
        if player_id in eligible_ids:
            by_player.setdefault(player_id, []).append(split)
    result = []
    for player_id, player_splits in by_player.items():
        if len(player_splits) > 1:
            totals = [split for split in player_splits
                      if split.get("isTotal") or not split.get("team")]
            if len(totals) == 1:
                player_splits = totals
            else:
                payload = api_get(f"/people/{player_id}/stats", {
                    "stats": "byDateRange", "group": group, "season": season,
                    "gameType": "R", "sportIds": 1,
                    "startDate": date.fromisoformat(start).strftime("%m/%d/%Y"),
                    "endDate": date.fromisoformat(through).strftime("%m/%d/%Y"),
                })
                player_splits = [split for stats in payload.get("stats", [])
                                 for split in stats.get("splits", [])]
                totals = [split for split in player_splits
                          if split.get("isTotal") or not split.get("team")]
                if len(totals) == 1:
                    player_splits = totals
        if len(player_splits) != 1:
            raise ValueError(f"Cannot determine combined {group} totals for player {player_id}")
        result.append(player_splits[0] | {"player": {"id": player_id}})
    return result


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def excel_value(field: str, value):
    if value in (None, "", ".---", "-.--"):
        return None
    if field in {"birth_date", "mlb_debut_date"}:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    if isinstance(value, str):
        try:
            return float(value) if "." in value else int(value)
        except ValueError:
            return value
    return value


def write_excel(path: Path, hitter_rows: list[dict], pitcher_rows: list[dict],
                period: tuple[str, str, str] | None = None) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise SystemExit("Excel output requires openpyxl. Install it with: python -m pip install openpyxl") from error

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows, fields in (
        ("Batters", hitter_rows, BIO_FIELDS + list(HITTING_FIELDS)),
        ("Pitchers", pitcher_rows, BIO_FIELDS + list(PITCHING_FIELDS)),
    ):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(fields)
        for row in rows:
            sheet.append([excel_value(field, row.get(field)) for field in fields])
        header = sheet[1]
        for cell in header:
            cell.fill = PatternFill("solid", fgColor="0F172A")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        if rows:
            table = Table(displayName=f"{sheet_name}Table", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
                showRowStripes=True, showColumnStripes=False,
            )
            sheet.add_table(table)
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 28)
            sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 10)
        for field in ("birth_date", "mlb_debut_date"):
            column = fields.index(field) + 1
            for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
                for cell in cells:
                    cell.number_format = "yyyy-mm-dd"
    workbook.properties.title = "MLB playoff-pool regular-season statistics"
    if period:
        start, through, roster_date = period
        workbook.properties.title += f" from {start} through {through}"
        workbook.properties.description = f"Current organization and roster status as of {roster_date}"
    workbook.properties.subject = "Official MLB Stats API regular-season totals"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_site_data(path: Path, hitter_rows: list[dict], pitcher_rows: list[dict],
                    season: int, start: str, through: str, roster_date: str) -> None:
    """Publish a range alongside previously exported ranges for this season."""
    payload = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"ranges": []}
    range_id = "ytd" if start == f"{season}-01-01" else f"{start}_{through}"
    fields = {
        "batters": {"G": "games_played", "AB": "at_bats", "H": "hits", "R": "runs",
                    "HR": "home_runs", "RBI": "rbi", "SB": "stolen_bases", "BB": "walks", "AVG": "batting_average"},
        "pitchers": {"G": "games", "IP": "innings_pitched", "W": "wins", "L": "losses",
                     "SV": "saves", "K": "strikeouts", "ERA": "era", "WHIP": "whip"},
    }
    def players(rows, group):
        return [{"id": row["player_id"], "name": row["player_name"],
                 "team": row["current_mlb_organization"],
                 "position": row.get("primary_position", ""),
                 "stats": {key: row.get(column, "") for key, column in fields[group].items()}}
                for row in rows]
    entry = {"id": range_id, "season": season, "start": start, "through": through,
             "roster_date": roster_date,
             "label": f"{season} year to date" if range_id == "ytd" else f"{start} – {through}",
             "teams": sorted(PLAYOFF_TEAMS),
             "batters": players(hitter_rows, "batters"), "pitchers": players(pitcher_rows, "pitchers")}
    ranges = [row for row in payload.get("ranges", [])
              if row["season"] == season and row["id"] != range_id]
    ranges.append(entry)
    ranges.sort(key=lambda row: (row["id"] != "ytd", row["start"], row["through"]))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"ranges": ranges}, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--season", type=int, default=date.today().year)
    parser.add_argument("--site", action="store_true", help="also publish this range to docs/data/player-stats.json")
    parser.add_argument("--start", help="first included date, YYYY-MM-DD (default: January 1 of --season)")
    parser.add_argument("--through", default=date.today().isoformat(), help="YYYY-MM-DD")
    parser.add_argument("--format", choices=("csv", "xlsx", "both"), default="both",
                        help="output two CSVs, one Excel workbook, or both (default: both)")
    args = parser.parse_args()

    try:
        through_date = date.fromisoformat(args.through)
        start_date = date.fromisoformat(args.start or f"{args.season}-01-01")
    except ValueError as error:
        parser.error(f"--start and --through must be YYYY-MM-DD: {error}")
    if through_date.year != args.season or start_date.year != args.season:
        parser.error("--start and --through must fall within --season")
    if start_date > through_date:
        parser.error("--start must be on or before --through")
    start = start_date.isoformat()
    through = through_date.isoformat()
    roster_date = date.today()

    affiliations, _ = current_affiliations(roster_date.isoformat(), roster_date.year, PLAYOFF_TEAMS)
    hitting = combined_splits(
        season_splits("hitting", args.season, through, start), "hitting",
        args.season, start, through, set(affiliations),
    )
    pitching = combined_splits(
        season_splits("pitching", args.season, through, start), "pitching",
        args.season, start, through, set(affiliations),
    )
    stat_ids = {row["player"]["id"] for row in hitting + pitching}
    people = people_details(stat_ids & affiliations.keys())

    hitter_rows = build_rows(hitting, HITTING_FIELDS, affiliations, people, "hitting")
    pitcher_rows = build_rows(pitching, PITCHING_FIELDS, affiliations, people, "pitching")
    output_dir = ROOT / "data" / str(args.season)
    period_label = f"{start}_through_{through}" if args.start else f"ytd_{through}"
    hitter_path = output_dir / f"mlb_hitters_{period_label}.csv"
    pitcher_path = output_dir / f"mlb_pitchers_{period_label}.csv"
    workbook_path = output_dir / f"mlb_players_{period_label}.xlsx"
    print(f"Regular-season statistics: {start} through {through} (inclusive).")
    print(f"Organization and roster status as of {roster_date.isoformat()}; stats include all MLB teams.")
    if args.format in {"csv", "both"}:
        write_csv(hitter_path, hitter_rows, BIO_FIELDS + list(HITTING_FIELDS))
        write_csv(pitcher_path, pitcher_rows, BIO_FIELDS + list(PITCHING_FIELDS))
        print(f"Wrote {len(hitter_rows)} batters to {hitter_path.relative_to(ROOT)}")
        print(f"Wrote {len(pitcher_rows)} pitchers to {pitcher_path.relative_to(ROOT)}")
    if args.format in {"xlsx", "both"}:
        write_excel(workbook_path, hitter_rows, pitcher_rows, (start, through, roster_date.isoformat()))
        print(f"Wrote Batters and Pitchers worksheets to {workbook_path.relative_to(ROOT)}")
    if args.site:
        site_path = ROOT / "docs/data/player-stats.json"
        write_site_data(site_path, hitter_rows, pitcher_rows, args.season, start, through, roster_date.isoformat())
        print(f"Updated {site_path.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
